import openpyxl, json, os

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIFOB_FILE = os.path.join(ROOT, 'DIFOB.xlsx')
OUT_FILE = os.path.join(ROOT, 'src', 'data', 'f06_data.json')

wb = openpyxl.load_workbook(DIFOB_FILE, data_only=True)
ws = [s for s in wb.worksheets if 'F-06' in s.title][0]

def to_float(v):
    if v is None: return 0.0
    if isinstance(v, (int, float)): return float(v)
    if isinstance(v, str):
        s = v.replace('R$', '').replace(' ', '').replace('.', '').replace(',', '.')
        try: return float(s)
        except: return 0.0
    return 0.0

def padronizar_mes(m):
    m_clean = str(m).strip().lower().replace('*', '')
    if 'jan' in m_clean: return 'Jan'
    if 'fev' in m_clean: return 'Fev'
    if 'mar' in m_clean: return 'Mar'
    if 'abr' in m_clean: return 'Abr'
    if 'mai' in m_clean: return 'Mai'
    if 'jun' in m_clean: return 'Jun'
    if 'jul' in m_clean: return 'Jul'
    if 'ago' in m_clean: return 'Ago'
    if 'set' in m_clean: return 'Set'
    if 'out' in m_clean: return 'Out'
    if 'nov' in m_clean: return 'Nov'
    if 'dez' in m_clean: return 'Dez'
    if '13' in m_clean: return '13º'
    return m

def extrair_linhas(r_ini, r_fim, ano):
    linhas = []
    for r in range(r_ini, r_fim + 1):
        m = ws.cell(r, 1).value
        desc = ws.cell(r, 2).value
        proc = ws.cell(r, 3).value
        fundeb = to_float(ws.cell(r, 4).value)
        fundeb_tes = to_float(ws.cell(r, 5).value)
        tesouro = to_float(ws.cell(r, 6).value)
        encargos = to_float(ws.cell(r, 7).value)
        tot_tes = to_float(ws.cell(r, 8).value)
        
        if not m and not desc: continue
        
        mes_p = padronizar_mes(m)
        tot_geral = fundeb + fundeb_tes + tesouro + encargos
        
        linhas.append({
            'ano': ano,
            'mes_raw': str(m).strip(),
            'mes': mes_p,
            'tipo_folha': str(desc).strip(),
            'processo': str(proc).strip() if proc else '',
            'fundeb': fundeb,
            'fundeb_tesouro': fundeb_tes,
            'tesouro': tesouro,
            'encargos': encargos,
            'total_tesouro': tot_tes,
            'total_geral': tot_geral,
            'projetado': 'PROJE' in str(proc).upper() or '*' in str(m)
        })
    return linhas

l2025 = extrair_linhas(6, 34, 2025)
l2026 = extrair_linhas(38, 67, 2026)

def agrupar_por_mes(linhas):
    meses_ordem = ['Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun', 'Jul', 'Ago', 'Set', 'Out', 'Nov', 'Dez', '13º']
    por_mes = {m: {'mes': m, 'fundeb': 0.0, 'fundeb_tesouro': 0.0, 'tesouro': 0.0, 'encargos': 0.0, 'total': 0.0, 'projetado': False, 'folhas': []} for m in meses_ordem}
    
    for item in linhas:
        m = item['mes']
        if m in por_mes:
            por_mes[m]['fundeb'] += item['fundeb']
            por_mes[m]['fundeb_tesouro'] += item['fundeb_tesouro']
            por_mes[m]['tesouro'] += item['tesouro']
            por_mes[m]['encargos'] += item['encargos']
            por_mes[m]['total'] += item['total_geral']
            if item['projetado']: por_mes[m]['projetado'] = True
            por_mes[m]['folhas'].append(item)
            
    return [por_mes[m] for m in meses_ordem if por_mes[m]['total'] > 0 or por_mes[m]['projetado']]

mensal_2025 = agrupar_por_mes(l2025)
mensal_2026 = agrupar_por_mes(l2026)

# Totais
tot_2025 = {
    'total': sum(x['total_geral'] for x in l2025),
    'fundeb': sum(x['fundeb'] for x in l2025),
    'fundeb_tesouro': sum(x['fundeb_tesouro'] for x in l2025),
    'tesouro': sum(x['tesouro'] for x in l2025),
    'encargos': sum(x['encargos'] for x in l2025),
}
tot_2026 = {
    'total': sum(x['total_geral'] for x in l2026),
    'fundeb': sum(x['fundeb'] for x in l2026),
    'fundeb_tesouro': sum(x['fundeb_tesouro'] for x in l2026),
    'tesouro': sum(x['tesouro'] for x in l2026),
    'encargos': sum(x['encargos'] for x in l2026),
    'realizado': sum(x['total_geral'] for x in l2026 if not x['projetado']),
    'projetado': sum(x['total_geral'] for x in l2026 if x['projetado']),
}

# Folhas por tipo
por_tipo_2026 = {}
for item in l2026:
    t = item['tipo_folha'].upper()
    por_tipo_2026[t] = por_tipo_2026.get(t, 0.0) + item['total_geral']

out_data = {
    'totais_2025': tot_2025,
    'totais_2026': tot_2026,
    'mensal_2025': mensal_2025,
    'mensal_2026': mensal_2026,
    'linhas_2025': l2025,
    'linhas_2026': l2026,
    'por_tipo_folha_2026': sorted([[k, v] for k, v in por_tipo_2026.items()], key=lambda x: -x[1]),
    'fonte': 'CFOP / SAGEP - Sistemas SIGRH e PAE',
    'nota': 'Ano 2026 composto por realizado (Jan a Out) e projeção CFOP (Nov, Dez e 13º Salário).'
}

with open(OUT_FILE, 'w', encoding='utf-8') as f:
    json.dump(out_data, f, ensure_ascii=False, indent=2)

print('Processamento da F-06 Folha de Pagamento concluído!')
print(f'Salvo em {OUT_FILE}')
print(f'Totais: 2025 = {tot_2025["total"]:,.2f} | 2026 = {tot_2026["total"]:,.2f} (Realizado: {tot_2026["realizado"]:,.2f}, Projetado: {tot_2026["projetado"]:,.2f})')
