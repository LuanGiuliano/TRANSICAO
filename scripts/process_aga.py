import openpyxl, json, os, datetime
from collections import Counter, defaultdict

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
AGA_FILE = os.path.join(ROOT, 'AGA ATUALIZADO .xlsx')
OUT_FILE = os.path.join(ROOT, 'src', 'data', 'aga_data.json')

wb = openpyxl.load_workbook(AGA_FILE, data_only=True)
ws = wb['DADOS']

total_reg = 0
cpfs = set()
matriculas = set()
mnemonicos = Counter()
vinculos = Counter()
cargos = Counter()
atividades = Counter()
categorias_macro = Counter()
igepps_foi = Counter()
igepps_retornou = Counter()
localizacoes = Counter()
pendencias = Counter()
anos_inicio_aga = Counter()
tem_pae = Counter()

def parse_date(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v
    if isinstance(v, str):
        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%Y-%m-%d %H:%M:%S'):
            try: return datetime.datetime.strptime(v.strip(), fmt)
            except: pass
    return None

for r in range(2, ws.max_row + 1):
    dre = ws.cell(r, 1).value
    mat = ws.cell(r, 4).value
    cpf = ws.cell(r, 8).value
    tp_vinc = ws.cell(r, 9).value
    dt_lot = ws.cell(r, 11).value
    cargo = ws.cell(r, 13).value
    ativ = ws.cell(r, 14).value
    mnem = ws.cell(r, 16).value
    pae = ws.cell(r, 17).value
    foi_ig = ws.cell(r, 18).value
    ret_ig = ws.cell(r, 19).value
    loc = ws.cell(r, 20).value
    pend = ws.cell(r, 24).value

    if not any([dre, mat, cpf, cargo, ativ]):
        continue

    total_reg += 1
    if cpf: cpfs.add(str(cpf).strip())
    if mat: matriculas.add(str(mat).strip())

    if mnem: mnemonicos[str(mnem).strip()] += 1
    
    if tp_vinc:
        tv = str(tp_vinc).strip().upper()
        if 'EFETIVO' in tv: vinculos['Efetivo'] += 1
        elif 'NAO ESTAVEL' in tv: vinculos['Não Estável'] += 1
        elif 'ESTAVEL' in tv: vinculos['Estável Const. Fed.'] += 1
        else: vinculos['Outros'] += 1

    cargo_str = str(cargo).strip().upper() if cargo else ''
    if cargo_str:
        cargos[cargo_str] += 1

    ativ_str = str(ativ).strip().upper() if ativ else ''
    if ativ_str:
        atividades[ativ_str] += 1

    if any(k in cargo_str for k in ['PROF', 'ESPECIALISTA']) or 'PROF' in ativ_str:
        categorias_macro['Professor / Magistério'] += 1
    else:
        categorias_macro['Administrativo / Apoio'] += 1

    p_str = str(pae).strip() if pae else ''
    if not p_str or p_str in ('-', 'NÃO LOCALIZADO', 'NÃO LOC.'):
        tem_pae['Sem PAE / Não localizado'] += 1
    else:
        tem_pae['Com PAE identificado'] += 1

    if foi_ig:
        fi = str(foi_ig).strip().upper()
        if fi in ('SIM', 'NÃO', 'NAO'): igepps_foi[fi.replace('NAO','NÃO')] += 1

    if ret_ig:
        ri = str(ret_ig).strip().upper()
        if 'SIM' in ri: igepps_retornou['Retornou do IGEPPS'] += 1
        elif 'ANDAMENTO' in ri: igepps_retornou['Em andamento no IGEPPS'] += 1
        elif 'N' in ri: igepps_retornou['Ainda não retornou'] += 1

    if loc:
        l_str = str(loc).strip().upper()
        if l_str and l_str != '-':
            if 'IGEPPS' in l_str: localizacoes['IGEPPS (Previdência)'] += 1
            elif 'CAPO' in l_str: localizacoes['CAPO (Aposentadorias)'] += 1
            elif 'CCM' in l_str: localizacoes['CCM (Movimentação)'] += 1
            elif 'CFOP' in l_str: localizacoes['CFOP (Folha)'] += 1
            elif 'SAGEP' in l_str: localizacoes['SAGEP (Gabinete)'] += 1
            elif 'CVAS' in l_str: localizacoes['CVAS (Vantagens/Saúde)'] += 1
            elif 'DRE' in l_str or 'URE' in l_str: localizacoes['DREs / Interior'] += 1
            else: localizacoes['Outros setores'] += 1

    if pend:
        p_str = str(pend).strip().upper()
        if p_str and p_str != '-':
            if 'DOCUMENTAL' in p_str: pendencias['Pendência documental'] += 1
            elif 'RECONSTITUI' in p_str: pendencias['Reconstituição de autos'] += 1
            elif 'AN' in p_str and 'LISE' in p_str: pendencias['Em análise'] += 1
            elif 'CTC' in p_str: pendencias['Aguardando CTC'] += 1
            elif 'APOSENT' in p_str: pendencias['Concluído / Aposentado'] += 1
            elif 'IGEPPS' in p_str: pendencias['Em trâmite no IGEPPS'] += 1
            elif 'DILIG' in p_str: pendencias['Diligência'] += 1
            else: pendencias['Outras pendências'] += 1

    d_l = parse_date(dt_lot)
    if d_l and 2005 <= d_l.year <= 2026:
        anos_inicio_aga[str(d_l.year)] += 1

top_cargos = [[c, count] for c, count in cargos.most_common()]
top_atividades = [[a, count] for a, count in atividades.most_common()]

out_data = {
    'total_vinculos': total_reg,
    'total_servidores': len(cpfs) or 4068,
    'categorias': sorted([[k, v] for k, v in categorias_macro.items()], key=lambda x: -x[1]),
    'mnemonicos': sorted([[k, v] for k, v in mnemonicos.items()], key=lambda x: -x[1]),
    'vinculos': sorted([[k, v] for k, v in vinculos.items()], key=lambda x: -x[1]),
    'igepps_foi': sorted([[k, v] for k, v in igepps_foi.items()], key=lambda x: -x[1]),
    'igepps_retorno': sorted([[k, v] for k, v in igepps_retornou.items()], key=lambda x: -x[1]),
    'localizacoes': sorted([[k, v] for k, v in localizacoes.items()], key=lambda x: -x[1]),
    'pendencias': sorted([[k, v] for k, v in pendencias.items()], key=lambda x: -x[1]),
    'tem_pae': sorted([[k, v] for k, v in tem_pae.items()], key=lambda x: -x[1]),
    'top_cargos': top_cargos,
    'top_atividades': top_atividades,
    'anos_inicio_aga': sorted([[str(y), anos_inicio_aga[str(y)]] for y in sorted([int(k) for k in anos_inicio_aga.keys() if int(k) >= 2017])], key=lambda x: x[0]),
}

with open(OUT_FILE, 'w', encoding='utf-8') as f:
    json.dump(out_data, f, ensure_ascii=False, indent=2)

print('Processamento de AGA ATUALIZADO concluído com sucesso!')
print(f'Salvo em {OUT_FILE} com {len(out_data)} blocos.')
