"""Extrai dados das planilhas DIOP/DIFOB/DIPSE para src/data/planilhas.json
e as imagens da F-07 para public/f07/. Rode: python scripts/extract.py"""
import json, re, zipfile, datetime, os
from collections import defaultdict
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'src', 'data')
os.makedirs(OUT, exist_ok=True)


def num(v):
    return float(v) if isinstance(v, (int, float)) else None


def txt(v):
    if v is None:
        return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%d/%m/%Y')
    s = str(v).strip()
    return s or None


def capa(ws):
    d = {}
    for r in ws.iter_rows(min_row=4, max_row=30, max_col=2):
        k, v = txt(r[0].value), txt(r[1].value)
        if k and v:
            d[k] = v
    return d


def top(d, n=12):
    return sorted(([k, round(v, 1)] for k, v in d.items()), key=lambda x: -x[1])[:n]


res = {}

# ---------------- DIOP ----------------
wb = openpyxl.load_workbook(os.path.join(ROOT, 'DIOP.xlsx'), data_only=True)
ws = {s.title.strip(): s for s in wb.worksheets}
res['DIOP'] = {'capa': capa(wb.worksheets[0])}

f1 = [s for t, s in ws.items() if t.startswith('F-01')][0]
sintese, link = [], None
for r in range(8, 30):
    a, b = f1.cell(r, 1).value, f1.cell(r, 2).value
    if isinstance(b, (int, float)) and a:
        sintese.append([txt(a), b])
    elif a and isinstance(b, str) and b.startswith('http'):
        link = b
total_geral = next((v for k, v in sintese if k.lower().startswith('total')), None)
sintese = [x for x in sintese if not x[0].lower().startswith('total')]

dre_tot, dre_vinc, cargo_tot, escolas, vinc_tot = (defaultdict(float), defaultdict(lambda: defaultdict(float)),
                                                  defaultdict(float), set(), defaultdict(float))
dre = escola = vinc = None
for row in f1.iter_rows(min_row=5, min_col=4, max_col=8, values_only=True):
    d, e, f, g, h = row
    if d and 'Total' not in str(d):
        dre = txt(d)
    if e and 'Total' not in str(e):
        escola = txt(e)
    if f and 'Total' not in str(f):
        vinc = txt(f)
    if g and isinstance(h, (int, float)):
        if dre and escola and vinc:
            dre_tot[dre] += h
            dre_vinc[dre][vinc] += h
            cargo_tot[txt(g)] += h
            vinc_tot[vinc] += h
            escolas.add((dre, escola))
res['DIOP']['F-01'] = {
    'sintese': sintese, 'total': total_geral, 'link': link,
    'por_dre': top(dre_tot, 40),
    'por_dre_vinculo': {k: {a: b for a, b in v.items()} for k, v in dre_vinc.items()},
    'cargos_top': top(cargo_tot, 12),
    'n_escolas': len(escolas), 'n_dres': len(dre_tot),
}

f2 = [s for t, s in ws.items() if t.startswith('F-02')][0]
comp = []
for r in range(5, f2.max_row + 1):
    a, b = f2.cell(r, 1).value, f2.cell(r, 2).value
    if a and isinstance(b, (int, float)) and 'Total' not in str(a):
        comp.append([txt(a), b])
res['DIOP']['F-01']['componentes'] = sorted(comp, key=lambda x: -x[1])[:25]
res['DIOP']['F-01']['n_componentes'] = len(comp)

# F-02: INTERIOR e RMB (demanda descoberta por modalidade)
mods = {'EDUCACAO ESPECIAL': 'Educação especial', 'FUND I': 'Fundamental I',
        'FUND I INTEGRAL': 'Fund. I integral', 'INTEGRAL': 'Integral', 'REG E EJA': 'Regular e EJA'}
f2d = {}
for regiao in ('INTERIOR', 'RMB'):
    sh = ws[regiao]
    blocos = {}
    for c in range(1, sh.max_column + 1):
        lab = txt(sh.cell(1, c).value)
        if lab in mods:
            blocos[lab] = c
    out = {}
    for lab, c0 in blocos.items():
        ch, qp, d_ch, disc_ch, total_info = 0.0, 0.0, defaultdict(float), defaultdict(float), None
        d_qp = defaultdict(float)
        for r in range(3, sh.max_row + 1):
            cl = sh.cell(r, c0).value
            if txt(cl) != regiao:
                continue
            d_, disc = txt(sh.cell(r, c0 + 1).value), txt(sh.cell(r, c0 + 3).value)
            c_h, q = num(sh.cell(r, c0 + 4).value) or 0, num(sh.cell(r, c0 + 5).value) or 0
            ch += c_h; qp += q
            d_ch[d_] += c_h; d_qp[d_] += q; disc_ch[disc] += c_h
        for r in range(3, 10):
            t = num(sh.cell(r, c0 + 6).value)
            if t is not None:
                total_info = t; break
        if ch or qp:
            out[mods[lab]] = {'ch': ch, 'prof': qp, 'total_planilha': total_info,
                              'dres': top(d_ch, 10), 'disciplinas': top(disc_ch, 10)}
    f2d[regiao] = out
res['DIOP']['F-02'] = f2d

# F-07: links + imagens
f7 = [s for t, s in ws.items() if t.startswith('F-07')][0]
links = []
for r in range(1, f7.max_row + 1):
    a = txt(f7.cell(r, 1).value)
    if a and a.startswith('http') and links:
        links[-1]['url'] = a
    elif a:
        links.append({'titulo': a, 'url': None})
res['DIOP']['F-07'] = {'links': [l for l in links if l['url']]}

z = zipfile.ZipFile(os.path.join(ROOT, 'DIOP.xlsx'))
pub = os.path.join(ROOT, 'public', 'f07')
os.makedirs(pub, exist_ok=True)
sheet_idx = [i for i, s in enumerate(wb.worksheets) if s.title.strip().startswith('F-07')][0] + 1
rels = z.read(f'xl/worksheets/_rels/sheet{sheet_idx}.xml.rels').decode()
dr = re.search(r'drawings/(drawing\d+)\.xml', rels).group(1)
drrels = z.read(f'xl/drawings/_rels/{dr}.xml.rels').decode()
imgs = []
for m in re.finditer(r'Target="\.\./media/([^"]+)"', drrels):
    name = m.group(1)
    if name.lower().endswith(('.png', '.jpg', '.jpeg')):
        with open(os.path.join(pub, name), 'wb') as fh:
            fh.write(z.read('xl/media/' + name))
        imgs.append('/f07/' + name)
res['DIOP']['F-07']['imagens'] = imgs

# ---------------- DIFOB ----------------
wb = openpyxl.load_workbook(os.path.join(ROOT, 'DIFOB.xlsx'), data_only=True)
res['DIFOB'] = {'capa': capa(wb.worksheets[0])}
f5 = [s for s in wb.worksheets if 'F-05' in s.title][0]
tab, link = [], None
cols = [txt(f5.cell(4, c).value) for c in range(1, 9)]
for r in range(5, 20):
    nome = txt(f5.cell(r, 1).value)
    if not nome or nome.startswith('*'):
        continue
    row = {'produto': nome}
    for c in range(2, 8):
        row[cols[c - 1]] = f5.cell(r, c).value
    l = txt(f5.cell(r, 8).value)
    if l:
        link = l
    tab.append(row)
nota = txt(f5.cell(7, 1).value)
res['DIFOB']['F-05'] = {'colunas': cols[1:7], 'linhas': tab, 'link': link, 'nota': nota}
res['DIFOB']['F-04'] = {}; res['DIFOB']['F-06'] = {}; res['DIFOB']['F-08'] = {}

# ---------------- DIPSE ----------------
wb = openpyxl.load_workbook(os.path.join(ROOT, 'DIPSE.xlsx'), data_only=True)
res['DIPSE'] = {'capa': capa(wb.worksheets[0])}
s = wb.worksheets[1]
pss = []
for r in range(9, 16):
    n = txt(s.cell(r, 1).value)
    if n and not n.startswith('Produto'):
        pss.append({'nome': n, 'vigencia': txt(s.cell(r, 2).value), 'modalidade': txt(s.cell(r, 3).value),
                    'link': txt(s.cell(r, 4).value)})
cab = [txt(s.cell(20, c).value) or '' for c in range(3, 8)]
cont = []
for r in range(21, 24):
    cont.append({'categoria': txt(s.cell(r, 1).value), 'total': num(s.cell(r, 2).value),
                 'desligamentos': [num(s.cell(r, c).value) or 0 for c in range(3, 8)]})
conc = {'nome': txt(s.cell(28, 1).value), 'validade': txt(s.cell(28, 2).value), 'obs': txt(s.cell(28, 5).value)}
cargos = []
for r in range(37, s.max_row + 1):
    cat, nome = txt(s.cell(r, 1).value), txt(s.cell(r, 4).value)
    cr, oc, vg = num(s.cell(r, 6).value), num(s.cell(r, 7).value), num(s.cell(r, 8).value)
    if cat and nome and cr is not None:
        cargos.append({'categoria': cat, 'cargo': re.sub(r'\s+', ' ', nome), 'natureza': txt(s.cell(r, 5).value),
                       'criado': cr, 'ocupado': oc or 0, 'vago': vg or 0})
res['DIPSE']['F-03'] = {'pss': pss, 'contratos': cont,
                        'contratos_cab': ['Dez/2026', 'Jan/2027', 'Fev/2027', 'Mar/2027', 'A partir de abr/2027'],
                        'concurso': conc, 'cargos': cargos}

with open(os.path.join(OUT, 'planilhas.json'), 'w', encoding='utf-8') as fh:
    json.dump(res, fh, ensure_ascii=False, indent=1)
print('OK', os.path.getsize(os.path.join(OUT, 'planilhas.json')), 'bytes; imagens F-07:', imgs)
print('F-01 total', total_geral, 'soma por_dre', sum(v for _, v in res['DIOP']['F-01']['por_dre']))
print('F-02', {r: {m: (d['ch'], d['prof'], d['total_planilha']) for m, d in v.items()} for r, v in f2d.items()})
print('cargos', len(cargos), 'pss', len(pss))
